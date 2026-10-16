"""Module 2 Excel practice workbook: creation, answer key and grading.

The practice workbook contains raw data, assumptions and yellow input cells
where the student writes formulas. ``grade_workbook`` reads the student's
file with openpyxl and evaluates each formula with the built-in engine, so
grading works even if the file was saved by a tool that does not store
calculated values.
"""
from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..core import metrics as m
from ..core.checker import format_value, within_tolerance
from ..data import generator as gen
from .formula_engine import Engine, ExcelError, display, functions_used, has_absolute_ref, shift_formula

NAVY = "1F3A5F"
INPUT = "FFF2CC"
KEY = "E2F0D9"
HEAD = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor=INPUT)
KEY_FILL = PatternFill("solid", fgColor=KEY)
THIN = Side(style="thin", color="B7C4D1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FIRST, LAST = 2, 32                    # Raw Data rows for 31 January nights
RATE_CHANGE = 0.05
ELASTICITY = -1.2


@dataclass
class Task:
    id: int
    cell: str
    title: str
    instruction: str
    hint: str
    answer: str                         # model formula
    unit: str = "number"
    accepted: Optional[set[str]] = None  # at least one of these functions must appear
    needs_absolute: bool = False
    fill_range: Optional[tuple[int, int]] = None   # rows filled down (same column)
    skill: str = "Excel formulas"


TASKS: list[Task] = [
    Task(1, "D5", "Total rooms sold", "Total rooms sold in January (Raw Data column E).",
         "Use SUM over the whole column of rooms sold, rows 2 to 32.",
         "=SUM('Raw Data'!E2:E32)", accepted={"SUM"}),
    Task(2, "D6", "Total room revenue", "Total room revenue in January (Raw Data column F).",
         "Same idea as task 1, on the revenue column.", "=SUM('Raw Data'!F2:F32)", "currency",
         accepted={"SUM"}),
    Task(3, "D7", "Available room nights", "Total available room nights (Raw Data column D).",
         "Available rooms already exclude out-of-order rooms. Add them up.",
         "=SUM('Raw Data'!D2:D32)", accepted={"SUM"}),
    Task(4, "D8", "Occupancy %", "Occupancy = rooms sold / available room nights. Point to your "
         "answers in D5 and D7 rather than retyping numbers.",
         "Divide D5 by D7. Format the cell as a percentage.", "=D5/D7", "percent"),
    Task(5, "D9", "ADR", "ADR = room revenue / rooms sold. Wrap it in IFERROR so it shows 0 "
         "instead of #DIV/0! if nothing was sold.",
         "=IFERROR(revenue / rooms sold, 0)", "=IFERROR(D6/D5,0)", "currency",
         accepted={"IFERROR"}),
    Task(6, "D10", "RevPAR", "RevPAR = room revenue / available room nights.",
         "Revenue divided by available rooms, or Occupancy x ADR.", "=D6/D7", "currency"),
    Task(7, "D11", "Weekend room revenue", "Room revenue for Weekend nights only (Raw Data column C "
         "says Weekday or Weekend).",
         "SUMIFS(sum_range, criteria_range, \"Weekend\")",
         "=SUMIFS('Raw Data'!F2:F32,'Raw Data'!C2:C32,\"Weekend\")", "currency",
         accepted={"SUMIFS", "SUMIF", "SUMPRODUCT"}),
    Task(8, "D12", "Weekday rooms sold", "Rooms sold on Weekday nights only.",
         "SUMIFS on column E with the criterion \"Weekday\".",
         "=SUMIFS('Raw Data'!E2:E32,'Raw Data'!C2:C32,\"Weekday\")",
         accepted={"SUMIFS", "SUMIF", "SUMPRODUCT"}),
    Task(9, "D13", "Weekend ADR", "ADR for weekend nights = weekend revenue / weekend rooms sold.",
         "You already have weekend revenue in D11. You need a SUMIFS for weekend rooms sold.",
         "=D11/SUMIFS('Raw Data'!E2:E32,'Raw Data'!C2:C32,\"Weekend\")", "currency",
         accepted={"SUMIFS", "SUMIF", "SUMPRODUCT"}),
    Task(10, "D14", "Budget ADR lookup", "Look up the budget ADR for \"Jan\" from the Budget sheet.",
         "INDEX(Budget!C2:C13, MATCH(\"Jan\", Budget!A2:A13, 0)) works in every Excel version. "
         "In Excel 2021/365 you can also use XLOOKUP(\"Jan\", Budget!A2:A13, Budget!C2:C13).",
         "=INDEX(Budget!C2:C13,MATCH(\"Jan\",Budget!A2:A13,0))", "currency",
         accepted={"XLOOKUP", "INDEX", "VLOOKUP"}),
    Task(11, "D15", "ADR variance vs budget", "ADR variance = actual ADR - budget ADR.",
         "Actual minus budget. A negative answer means ADR was below budget.", "=D9-D14", "currency"),
    Task(12, "D16", "Budget flag", "Show \"Below budget\" if actual occupancy (D8) is below budget "
         "occupancy (Budget!B2), otherwise \"On or above budget\".",
         "IF(D8<Budget!B2, \"Below budget\", \"On or above budget\")",
         "=IF(D8<Budget!B2,\"Below budget\",\"On or above budget\")", "text", accepted={"IF"}),
    Task(13, "D17", "Busy nights", "Count the nights where at least 130 rooms were sold.",
         "COUNTIFS(range, \">=130\")", "=COUNTIFS('Raw Data'!E2:E32,\">=130\")",
         accepted={"COUNTIFS", "COUNTIF"}),
    Task(14, "D18", "Average weekday rooms sold", "Average rooms sold per Weekday night.",
         "AVERAGEIFS(average_range, criteria_range, criterion)",
         "=AVERAGEIFS('Raw Data'!E2:E32,'Raw Data'!C2:C32,\"Weekday\")",
         accepted={"AVERAGEIFS"}),
    Task(15, "D19", "Rate scenario revenue", "Projected January revenue if ADR changes by the rate "
         "change on the Assumptions sheet (B4) and demand responds with the elasticity in B5: "
         "Revenue x (1 + change) x (1 + change)^elasticity.",
         "Point to Assumptions!B4 and Assumptions!B5. The ^ operator raises to a power.",
         "=D6*(1+Assumptions!B4)*(1+Assumptions!B4)^Assumptions!B5", "currency"),
    Task(16, "I2", "Share of month revenue (fill down)",
         "In Raw Data!I2 calculate this night's share of January room revenue, then copy the "
         "formula down to I32. The total must stay fixed when you copy, so anchor it with $.",
         "=F2/SUM($F$2:$F$32). Without the $ signs the range slides down as you copy.",
         "=F2/SUM($F$2:$F$32)", "percent", needs_absolute=True, fill_range=(FIRST, LAST)),
]


# ----------------------------------------------------------- expected values

def january() -> pd.DataFrame:
    from ..data.loader import data
    d = data().daily
    return d[d["date"].str.startswith(f"{gen.YEAR}-01")].reset_index(drop=True)


def budget_table() -> pd.DataFrame:
    return gen.rooms_budget()


def expected_values(jan: pd.DataFrame | None = None, bud: pd.DataFrame | None = None) -> dict:
    """Answer key computed directly with pandas (independent of the engine)."""
    jan = january() if jan is None else jan
    bud = budget_table() if bud is None else bud
    sold, rev, avail = jan.rooms_sold.sum(), jan.room_revenue.sum(), jan.rooms_available.sum()
    wk = jan[jan.day_type == "Weekend"]
    wd = jan[jan.day_type == "Weekday"]
    adr_ = m.adr(rev, sold)
    b_adr = float(bud.loc[0, "budget_adr"])
    occ = m.occupancy(sold, avail)
    out = {
        1: sold, 2: rev, 3: avail, 4: occ, 5: adr_, 6: m.revpar(rev, avail),
        7: wk.room_revenue.sum(), 8: wd.rooms_sold.sum(),
        9: m.adr(wk.room_revenue.sum(), wk.rooms_sold.sum()),
        10: b_adr, 11: adr_ - b_adr,
        12: "Below budget" if occ < float(bud.loc[0, "budget_occupancy"]) else "On or above budget",
        13: int((jan.rooms_sold >= 130).sum()), 14: wd.rooms_sold.mean(),
        15: rev * (1 + RATE_CHANGE) * (1 + RATE_CHANGE) ** ELASTICITY,
        16: [r / rev for r in jan.room_revenue],
    }
    return out


# -------------------------------------------------------------- building

def xlsx_formula(f: str) -> str:
    """Functions newer than Excel 2010 must be stored with the _xlfn. prefix
    inside the file, otherwise Excel shows #NAME?. Excel adds it itself when
    a user types the formula; we must add it when writing with openpyxl."""
    return f.replace("XLOOKUP(", "_xlfn.XLOOKUP(")

def _header(ws, row, values, widths=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(path, answer_key: bool = False):
    jan, bud = january(), budget_table()
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Instructions"
    title = "ANSWER KEY - Module 2 Excel Practice" if answer_key else "Module 2 - Excel Practice Workbook"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    lines = [
        f"Hotel: {gen.HOTEL_NAME} (fictional), 150 rooms, January {gen.YEAR}.",
        gen.DISCLAIMER,
        "",
        "How to use this workbook",
        "1. Go to the 'Exercises' sheet. Yellow cells are yours: type a formula starting with '='.",
        "2. Use cell references, not typed numbers, so the model updates when data changes.",
        "3. Task 16 is on the 'Raw Data' sheet (column I): write one formula in I2 and copy it down.",
        "4. Save the file (keep the .xlsx format), then in the app open Excel Workbook Center",
        "   and click 'Grade my workbook'. You will get feedback on every task.",
        "",
        "Model structure used here (and in professional models):",
        "   Raw Data  ->  Assumptions  ->  Calculations (Exercises)  ->  Outputs",
        "Colour code: yellow = your input, blue headers = labels, green = answer key.",
        "Data validation: Assumptions!B4 only accepts -30% to +30%; B6 is a drop-down list.",
        "Conditional formatting: Raw Data rooms sold are colour-scaled (red = low, green = high).",
    ]
    for i, t in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=t)
    ws["A6"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 100

    rd = wb.create_sheet("Raw Data")
    _header(rd, 1, ["Date", "Weekday", "Day Type", "Rooms Available", "Rooms Sold",
                    "Room Revenue", "Cancellations", "Out of Order", "Share of Month Revenue"],
            [12, 12, 11, 15, 12, 15, 13, 12, 20])
    for i, r in jan.iterrows():
        row = FIRST + i
        vals = [dt.date.fromisoformat(r.date), r.weekday, r.day_type, int(r.rooms_available),
                int(r.rooms_sold), float(r.room_revenue), int(r.cancellations), int(r.rooms_out_of_order)]
        for c, v in enumerate(vals, 1):
            cell = rd.cell(row=row, column=c, value=v)
            cell.border = BOX
        rd.cell(row=row, column=1).number_format = "dd-mmm-yyyy"
        rd.cell(row=row, column=6).number_format = "#,##0.00"
        ci = rd.cell(row=row, column=9)
        ci.fill = KEY_FILL if answer_key else INPUT_FILL
        ci.number_format = "0.00%"
        ci.border = BOX
        if answer_key:
            ci.value = shift_formula(TASKS[15].answer, row - FIRST, 0)
    rd.conditional_formatting.add(f"E{FIRST}:E{LAST}", ColorScaleRule(
        start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
        mid_color="FFEB84", end_type="max", end_color="63BE7B"))
    rd.freeze_panes = "A2"

    a = wb.create_sheet("Assumptions")
    _header(a, 1, ["Assumption", "Value", "Note"], [34, 14, 60])
    rows = [("Rooms in inventory", 150, "Physical rooms"),
            ("Days in January", 31, ""),
            ("Rate change scenario", RATE_CHANGE, "Allowed range -30% to +30% (data validation)"),
            ("Price elasticity of demand", ELASTICITY, "-1.2: a 1% rate rise loses about 1.2% of rooms sold"),
            ("Scenario selected", "Base", "Drop-down: Base / Upside / Downside"),
            ]
    for i, (k, v, n) in enumerate(rows, 2):
        a.cell(row=i, column=1, value=k).border = BOX
        c = a.cell(row=i, column=2, value=v)
        c.border, c.fill = BOX, INPUT_FILL
        a.cell(row=i, column=3, value=n)
    a["B4"].number_format = "0.0%"
    dv = DataValidation(type="decimal", operator="between", formula1="-0.3", formula2="0.3",
                        showErrorMessage=True, errorTitle="Out of range",
                        error="Enter a rate change between -30% and +30%.")
    a.add_data_validation(dv)
    dv.add("B4")
    dv2 = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False)
    a.add_data_validation(dv2)
    dv2.add("B6")

    b = wb.create_sheet("Budget")
    _header(b, 1, ["Month", "Budget Occupancy", "Budget ADR", "Budget Room Revenue"], [10, 18, 14, 22])
    for i, r in bud.iterrows():
        vals = [r.month_name, float(r.budget_occupancy), float(r.budget_adr), float(r.budget_room_revenue)]
        for c, v in enumerate(vals, 1):
            b.cell(row=i + 2, column=c, value=v).border = BOX
        b.cell(row=i + 2, column=2).number_format = "0.0%"
        b.cell(row=i + 2, column=3).number_format = "#,##0.00"
        b.cell(row=i + 2, column=4).number_format = "#,##0"

    e = wb.create_sheet("Exercises", 1)
    e["A1"] = "Exercises - write a formula in each yellow cell"
    e["A1"].font = Font(bold=True, size=14, color=NAVY)
    e["A2"] = "Tip: hints are in column E. The app explains any mistakes when you grade the file."
    _header(e, 4, ["#", "Task", "Instruction", "Your formula", "Hint"], [5, 26, 70, 20, 60])
    for t in TASKS:
        if t.fill_range:
            continue
        r = int(t.cell[1:])
        e.cell(row=r, column=1, value=t.id).border = BOX
        e.cell(row=r, column=2, value=t.title).border = BOX
        ci = e.cell(row=r, column=3, value=t.instruction)
        ci.alignment = Alignment(wrap_text=True, vertical="top")
        ci.border = BOX
        d = e.cell(row=r, column=4)
        d.fill, d.border = (KEY_FILL if answer_key else INPUT_FILL), BOX
        d.number_format = {"percent": "0.00%", "currency": "#,##0.00"}.get(t.unit, "#,##0.00")
        if answer_key:
            d.value = xlsx_formula(t.answer)
        e.cell(row=r, column=5, value=t.hint).alignment = Alignment(wrap_text=True, vertical="top")
        e.row_dimensions[r].height = 45
    r = 21
    e.cell(row=r, column=2, value="Task 16").font = Font(bold=True)
    e.cell(row=r, column=3, value=TASKS[15].instruction).alignment = Alignment(wrap_text=True)
    e.row_dimensions[r].height = 45
    e.conditional_formatting.add("D16", CellIsRule(operator="equal", formula=['"Below budget"'],
                                                   font=Font(color="C00000", bold=True)))
    e.freeze_panes = "A5"

    if isinstance(path, (str, Path)):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --------------------------------------------------------------- grading

@dataclass
class TaskResult:
    task: Task
    status: str            # correct / incorrect / not_attempted / hardcoded / partial
    formula: str
    value_text: str
    expected_text: str
    feedback: str

    @property
    def correct(self) -> bool:
        return self.status == "correct"


def engine_from_workbook(wb) -> Engine:
    sheets = {}
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells[c.coordinate] = c.value
        sheets[ws.title] = cells
    return Engine(sheets, wb.worksheets[0].title)


def _compare(value, expected, unit) -> bool:
    if isinstance(value, ExcelError):
        return False
    if unit == "text":
        return str(value).strip().lower() == str(expected).strip().lower()
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    tol = {"percent": (0.00005, 0.0), "currency": (0.01, 0.0)}.get(unit, (0.01, 0.0))
    return within_tolerance(float(value), float(expected), *tol)


def grade_task(engine: Engine, cells_sheet: str, t: Task, expected) -> TaskResult:
    sheet = "Raw Data" if t.fill_range else cells_sheet
    if t.fill_range:
        col = t.cell[0]
        first = engine.sheets[sheet].get(f"{col}{t.fill_range[0]}")
        if first is None:
            return TaskResult(t, "not_attempted", "", "", "", "Cell I2 is empty.")
        wrong = []
        for i, row in enumerate(range(t.fill_range[0], t.fill_range[1] + 1)):
            raw = engine.sheets[sheet].get(f"{col}{row}")
            if raw is None:
                wrong.append(f"{col}{row} is empty")
                continue
            v = engine.cell_value(sheet, f"{col}{row}")
            if not _compare(v, expected[i], t.unit):
                wrong.append(f"{col}{row} = {display(v)}")
        f = str(first)
        if not wrong:
            if not f.startswith("="):
                return TaskResult(t, "hardcoded", f, "values typed", "", "The values are right but "
                                  "they were typed in. Use one formula and copy it down.")
            return TaskResult(t, "correct", f, "all 31 rows correct", "", "Correct. The $ signs kept "
                              "the total fixed while the row reference moved.")
        msg = f"{len(wrong)} of 31 cells are wrong (first: {wrong[0]})."
        if f.startswith("=") and not has_absolute_ref(f):
            msg += (" Your formula has no $ signs, so when you copied it down the SUM range slid "
                    "down too. Use SUM($F$2:$F$32) to lock it.")
        return TaskResult(t, "incorrect", f, "", "", msg)

    raw = engine.sheets.get(sheet, {}).get(t.cell)
    exp_text = format_value(expected, t.unit) if t.unit != "text" else str(expected)
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return TaskResult(t, "not_attempted", "", "", exp_text, "Not attempted yet.")
    f = str(raw)
    v = engine.cell_value(sheet, t.cell)
    vt = display(v)
    if not f.startswith("="):
        ok = _compare(v, expected, t.unit)
        return TaskResult(t, "hardcoded" if ok else "incorrect", f, vt, exp_text,
                          ("The number is right but it was typed in, not calculated. Use a formula "
                           "with cell references so the model updates when data changes.") if ok else
                          "This is a typed value, not a formula, and it is not correct.")
    if isinstance(v, ExcelError):
        why = v.detail or "Excel would show this error."
        return TaskResult(t, "incorrect", f, v.code, exp_text,
                          f"Your formula returns {v.code}. {why} Hint: {t.hint}")
    ok = _compare(v, expected, t.unit)
    used = functions_used(f)
    if ok and t.accepted and not (used & t.accepted):
        return TaskResult(t, "partial", f, vt, exp_text,
                          f"Right answer, but this task practises {' or '.join(sorted(t.accepted))}. "
                          "Try again using that function.")
    if ok:
        return TaskResult(t, "correct", f, vt, exp_text, "Correct.")
    fb = "Not correct."
    if t.unit == "percent" and isinstance(v, (int, float)) and abs(v - expected * 100) < 0.05:
        fb += " You multiplied by 100. Keep the decimal and format the cell as %."
    elif isinstance(v, (int, float)) and t.unit != "text" and expected and abs(v) > 0 and \
            abs(v / expected - 1) < 0.05:
        fb += " Very close: check that your range covers all rows 2 to 32."
    return TaskResult(t, "incorrect", f, vt, exp_text, f"{fb} Hint: {t.hint}")


def grade_workbook(path: Path) -> list[TaskResult]:
    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except Exception as exc:          # corrupt file, wrong format, file open elsewhere
        raise ValueError(f"Could not open the workbook: {exc}. Make sure it is a .xlsx file "
                         "and that it is closed in Excel.") from exc
    for needed in ("Exercises", "Raw Data", "Budget", "Assumptions"):
        if needed not in wb.sheetnames:
            raise ValueError(f"The workbook has no '{needed}' sheet. Grade the practice workbook "
                             "created by this app, and do not rename its sheets.")
    eng = engine_from_workbook(wb)
    exp = expected_values()
    return [grade_task(eng, "Exercises", t, exp[t.id]) for t in TASKS]


def evaluate_lab_formula(task: Task, formula: str, engine: Engine) -> TaskResult:
    """In-app Formula Lab: evaluate a typed formula as if entered in the task cell."""
    sheet = "Raw Data" if task.fill_range else "Exercises"
    formula = formula.strip()
    if formula and not formula.startswith("="):
        formula = "=" + formula
    if task.fill_range:
        for i, row in enumerate(range(task.fill_range[0], task.fill_range[1] + 1)):
            engine.sheets[sheet][f"I{row}"] = shift_formula(formula, i, 0) if formula else None
    else:
        engine.sheets[sheet][task.cell] = formula or None
    engine._cache.clear()
    exp = expected_values()
    return grade_task(engine, "Exercises", task, exp[task.id])


def lab_engine() -> Engine:
    """Engine pre-loaded with the practice workbook data plus the model answers
    of every task, so later tasks that reference earlier cells still work."""
    buf = io.BytesIO()
    build_workbook(buf, answer_key=True)
    buf.seek(0)
    return engine_from_workbook(openpyxl.load_workbook(buf))
