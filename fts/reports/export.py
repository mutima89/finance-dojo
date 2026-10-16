"""Exports: Excel tables, management report, progress/competency report and
the instructor answer key."""
from __future__ import annotations

import datetime as dt
import html
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from ..content import capstone, curriculum
from ..content.base import SKILLS
from ..core import fpa
from ..core import metrics as m
from ..core.checker import format_value
from ..data import generator as gen

DISCLAIMER_CERT = ("Completing this training simulator is not equivalent to a professional certification and does "
                   "not guarantee employment.")


def dataframe_to_xlsx(sheets: dict[str, pd.DataFrame], path, title: str = "") -> Path:
    path = Path(path)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], startrow=3)
            ws = xw.sheets[name[:31]]
            ws["A1"] = title or name
            ws["A1"].font = Font(bold=True, size=13, color="1F3A5F")
            ws["A2"] = gen.DISCLAIMER
            ws["A2"].font = Font(italic=True, size=8, color="666666")
            for cell in ws[4]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A5F")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 45)
            for row in ws.iter_rows(min_row=5):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.0%" if abs(c.value) < 2 and c.value != 0 and "%" in str(
                            ws.cell(row=4, column=c.column).value) else "#,##0.00"
    return path


def management_report_text(pnl: pd.DataFrame, ds, months: list[int], period: str) -> str:
    t = fpa.period_table(pnl, months)
    mr = ds.monthly_rooms[ds.monthly_rooms.month.isin(months)]
    rb = ds.rooms_budget[ds.rooms_budget.month.isin(months)]
    occ = m.occupancy(mr.sold_room_nights.sum(), mr.available_room_nights.sum())
    b_occ = m.occupancy(rb.budget_sold_room_nights.sum(), rb.available_room_nights.sum())
    adr = m.adr(mr.room_revenue.sum(), mr.sold_room_nights.sum())
    revpar = m.revpar(mr.room_revenue.sum(), mr.available_room_nights.sum())
    b_revpar = m.revpar(rb.budget_room_revenue.sum(), rb.available_room_nights.sum())
    rev, ebitda = t.loc["Total revenue"], t.loc["EBITDA"]
    top = fpa.top_unfavorable(t)

    def fm(x):
        return format_value(x, "currency", 0)

    lines = [
        f"{gen.HOTEL_NAME.upper()} - MANAGEMENT REPORT - {period} {gen.YEAR}",
        f"Prepared {dt.date.today():%d %b %Y} · fictional training data",
        "",
        "HEADLINE",
        f"Total revenue {fm(rev.Actual)} was {fm(abs(rev.Variance))} ({abs(rev['Variance %']) * 100:.1f}%) "
        f"{'below' if rev.Variance < 0 else 'above'} budget and EBITDA {fm(ebitda.Actual)} was "
        f"{fm(abs(ebitda.Variance))} {'below' if ebitda.Variance < 0 else 'above'} budget. "
        "[Edit: add the main reason in one clause.]",
        "",
        "KEY PERFORMANCE INDICATORS          Actual        Budget      Variance",
        f"Occupancy                       {occ * 100:9.1f}%    {b_occ * 100:9.1f}%   {(occ - b_occ) * 100:+8.1f} pts",
        f"RevPAR                          {revpar:10.2f}    {b_revpar:10.2f}   {revpar - b_revpar:+10.2f}",
        f"ADR                             {adr:10.2f}",
        f"Total revenue              {rev.Actual:15,.0f} {rev.Budget:13,.0f} {rev.Variance:+13,.0f}",
        f"GOP                        {t.at['Gross operating profit (GOP)', 'Actual']:15,.0f} "
        f"{t.at['Gross operating profit (GOP)', 'Budget']:13,.0f} {t.at['Gross operating profit (GOP)', 'Variance']:+13,.0f}",
        f"EBITDA                     {ebitda.Actual:15,.0f} {ebitda.Budget:13,.0f} {ebitda.Variance:+13,.0f}",
        f"EBITDA margin              {ebitda.Actual / rev.Actual * 100:14.1f}% {ebitda.Budget / rev.Budget * 100:12.1f}%",
        "",
        "THREE LARGEST UNFAVORABLE VARIANCES",
    ]
    for i, (k, r) in enumerate(top.iterrows(), 1):
        lines.append(f"{i}. {k}: {fm(r.Variance)} ({r['Variance %'] * 100:+.1f}%). Driver: [Edit: evidence]. "
                     "Action: [Edit: owner, date].")
    lines += ["", "OUTLOOK AND RISKS",
              "[Edit: revised forecast for the rest of the year, pace vs last year, key risks.]",
              "", "RECOMMENDED ACTIONS", "1. [Edit]", "2. [Edit]", "",
              "Assumptions and limitations: figures are unaudited management accounts (fictional training data); "
              "forecasts are estimates, not guarantees."]
    return "\n".join(lines)


def text_report_html(text: str, path, title: str) -> Path:
    path = Path(path)
    body = html.escape(text)
    path.write_text(f"""<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;max-width:900px;margin:30px auto;color:#1b2430}}
pre{{font-family:Consolas,Menlo,monospace;font-size:13px;white-space:pre-wrap;line-height:1.45;
background:#f6f8fb;border:1px solid #dde3ea;border-radius:8px;padding:18px}}
h1{{color:#1f3a5f}}</style></head><body><h1>{html.escape(title)}</h1><pre>{body}</pre>
<p style="color:#777;font-size:12px">{html.escape(gen.DISCLAIMER)}</p></body></html>""", encoding="utf-8")
    return path


# ------------------------------------------------------------ progress

def progress_frames(ctx_db, skills) -> dict[str, pd.DataFrame]:
    status = ctx_db.all_exercise_status()
    secs = ctx_db.lesson_seconds()
    tasks = ctx_db.all_tasks()
    from ..core import mastery
    rows = []
    for mod in curriculum.modules():
        for l in mod.lessons:
            best = ctx_db.best_quiz(l.id)
            rows.append({"Module": f"{mod.number}. {mod.title}", "Lesson": l.title,
                         "Progress": mastery.lesson_progress_fraction(l, status, best, tasks.get(l.task.id)),
                         "Complete": "Yes" if mastery.lesson_complete(l, status, best, tasks.get(l.task.id)) else "No",
                         "Best quiz": f"{best[0]}/{best[1]}" if best else "-",
                         "Exercises solved": sum(1 for e in l.exercises if status.get(e.id, {}).get("solved")),
                         "Exercises": len(l.exercises),
                         "Attempts": sum(status.get(e.id, {}).get("attempts", 0) for e in l.exercises),
                         "Solutions revealed": sum(1 for e in l.exercises if status.get(e.id, {}).get("revealed")),
                         "Practical task saved": "Yes" if (tasks.get(l.task.id) or {}).get("content") else "No",
                         "Active minutes": round(secs.get(l.id, 0) / 60, 1)})
    lessons = pd.DataFrame(rows)
    sk = pd.DataFrame([{"Skill": s.skill, "Score": s.score, "Level": s.level,
                        "Exercise component": round(s.exercise_score * 100, 1),
                        "Quiz component": round(s.quiz_score * 100, 1),
                        "Workbook component": "-" if s.workbook_score is None else round(s.workbook_score * 100, 1),
                        "Weak topic": "Yes" if s.weak else "", "Next requirement": s.reason} for s in skills])
    ex = pd.DataFrame([{"Exercise": e.id, "Skill": e.skill, "Practical": "Yes" if e.practical else "",
                        "Attempts": status.get(e.id, {}).get("attempts", 0),
                        "Solved on attempt": status.get(e.id, {}).get("solved_on") or "",
                        "Hints used": status.get(e.id, {}).get("hints_used", 0),
                        "Solution revealed": "Yes" if status.get(e.id, {}).get("revealed") else ""}
                       for _, _, e in curriculum.all_exercises()])
    wb = ctx_db.workbook_results()
    wbdf = pd.DataFrame([{"Task": i, "Result": wb.get(i, "not attempted")} for i in range(1, 17)])
    return {"Skills": sk, "Lessons": lessons, "Exercises": ex, "Excel tasks": wbdf}


def progress_report_html(ctx_db, skills, path) -> Path:
    path = Path(path)
    f = progress_frames(ctx_db, skills)
    name = ctx_db.setting("student_name") or "Student"
    best = ctx_db.best_capstone()
    n_att, n_ok = ctx_db.attempts_total()
    cap = f"{best['total']:.1f}/100 ({'passed' if best['passed'] else 'not yet passed'})" if best else "not attempted"
    tables = "".join(f"<h2>{html.escape(k)}</h2>" + v.to_html(index=False, float_format=lambda x: f"{x:,.2f}",
                                                                 border=0, classes="t")
                     for k, v in f.items())
    path.write_text(f"""<!doctype html><html><head><meta charset="utf-8"><title>Competency report</title>
<style>body{{font-family:Segoe UI,Arial,sans-serif;max-width:1100px;margin:30px auto;color:#1b2430}}
h1{{color:#1f3a5f}} table.t{{border-collapse:collapse;width:100%;font-size:13px;margin-bottom:20px}}
.t th{{background:#1f3a5f;color:#fff;padding:6px;text-align:left}} .t td{{border-bottom:1px solid #e3e7ed;padding:5px}}
.note{{background:#fff7e0;border:1px solid #f0d58a;padding:10px;border-radius:6px}}</style></head><body>
<h1>Final competency report: {html.escape(name)}</h1>
<p>Generated {dt.datetime.now():%d %b %Y %H:%M} · Finance Dojo</p>
<p><b>Exercise attempts:</b> {n_att} ({n_ok} correct) · <b>Capstone:</b> {cap}</p>
<p class="note">{DISCLAIMER_CERT} Levels: Beginner &lt; 40 ≤ Developing &lt; 70 ≤ Competent. "Job-ready for supervised
tasks" also requires a score of 85+, all practical exercises solved without revealing the solution, all practical
tasks saved and a passed capstone. Active time counts only while a lesson is open with recent keyboard or mouse activity.</p>
{tables}</body></html>""", encoding="utf-8")
    return path


def progress_report_xlsx(ctx_db, skills, path) -> Path:
    return dataframe_to_xlsx(progress_frames(ctx_db, skills), path, "Competency report")


# ------------------------------------------------------------ answer key

def answer_key_markdown() -> str:
    out = ["# Instructor answer key", "",
           "Finance Dojo. All data is fictional. Answers are generated from the same "
           "code the app uses to check students, and selected values are verified by independent hand calculations "
           "in the automated tests (`tests/test_content.py`).", "",
           "Tolerance: percentages ±0.06 percentage points; currency ±$0.60 or 0.05%; numbers ±0.01 or 0.05% unless "
           "noted. Equivalent formats (75%, 0.75, 75) are accepted.", ""]
    for mod in curriculum.modules():
        out += [f"## Module {mod.number}: {mod.title}", ""]
        for l in mod.lessons:
            out += [f"### {l.title}", ""]
            for kind, e in [("Guided", l.guided)] + [("Independent", x) for x in l.independent]:
                ans = (e.choices[e.expected] if e.unit == "choice"
                       else format_value(float(e.expected), e.unit, 3 if e.unit == "percent" else 2))
                out += [f"**{kind} `{e.id}`** ({e.skill}{', practical' if e.practical else ''})", "",
                        f"> {e.prompt}", "", f"- **Answer:** {ans}", f"- **Working:** {e.solution}"]
                if e.mistakes:
                    out.append("- **Common wrong answers:** " + "; ".join(
                        f"{format_value(mk.value, e.unit if e.unit != 'choice' else 'number', 2)}: {mk.explanation}"
                        for mk in e.mistakes))
                out.append("")
            out += ["**Quiz**", ""]
            for i, q in enumerate(l.quiz, 1):
                out.append(f"{i}. {q.question} → **{q.options[q.answer]}**. {q.explanation}")
            out += ["", f"**Practical task: {l.task.title}.** Marking checklist: " + "; ".join(l.task.checklist), ""]
    out += ["## Module 2 Excel workbook", "",
            "| # | Cell | Model formula | Expected value |", "|---|---|---|---|"]
    from ..excel import workbook as wbk
    exp = wbk.expected_values()
    for t in wbk.TASKS:
        val = "per row: F_n / SUM(F2:F32)" if t.fill_range else (
            exp[t.id] if t.unit == "text" else format_value(float(exp[t.id]), t.unit, 2))
        out.append(f"| {t.id} | {'Raw Data!' if t.fill_range else 'Exercises!'}{t.cell} | `{t.answer}` | {val} |")
    out += ["", "Accepted alternatives: XLOOKUP instead of INDEX/MATCH (task 10); SUMIF/SUMPRODUCT for SUMIFS tasks; "
            "COUNTIF for task 13.", "", "## Capstone", "",
            f"Pass mark {capstone.PASS_MARK}/100. Numeric and choice items are auto-scored; written items get an "
            "automatic structure score. Instructors should also judge quality with the rubric below.", "",
            "| Step | Item | Points | Expected |", "|---|---|---|---|"]
    for st in capstone.steps():
        for it in st.items:
            out.append(f"| {st.number}. {st.title} | {it.label} | {it.points:g} | {capstone.expected_display(it)} |")
    out += ["", "### Qualitative rubric for written work (instructor use)", "",
            "- **Executive recommendation (15):** recommendation first (3); supported by 3+ correct numbers (3); "
            "drivers explained with evidence (3); risks and mitigations (3); assumptions/limitations and next steps "
            "with owners (3).",
            "- **Variance explanation (5):** identifies room revenue (weekday corporate demand, OTA shift) and "
            "other operating expenses (commissions, utilities tariff, July chiller repair) with numbers.",
            "- **Forecast method (4):** states the ratio method, that it assumes Q4 under-performance persists into "
            "January, and that it ignores bookings on hand (pickup data would improve it).",
            "", "### Data issues built into the capstone raw file", "",
            f"- Duplicate row: {gen.CAPSTONE_ISSUES['duplicate_month']}",
            f"- Missing sold room nights: {gen.CAPSTONE_ISSUES['missing_sold_month']} (rebuild = revenue / reported ADR)",
            f"- Room revenue keyed in thousands: {gen.CAPSTONE_ISSUES['thousands_month']}",
            f"- Sold > available (transposed digits): {gen.CAPSTONE_ISSUES['impossible_occupancy_month']} "
            "(rebuild = revenue / reported ADR)", ""]
    return "\n".join(out)
