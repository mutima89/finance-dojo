import shutil
import subprocess

import openpyxl
import pytest

from fts.excel import workbook as wbk


@pytest.fixture(scope="module")
def files(tmp_path_factory):
    d = tmp_path_factory.mktemp("wb")
    return wbk.build_workbook(d / "practice.xlsx"), wbk.build_workbook(d / "key.xlsx", answer_key=True), d


def test_answer_key_grades_fully_correct(files):
    _, key, _ = files
    res = wbk.grade_workbook(key)
    bad = [(r.task.id, r.status, r.feedback) for r in res if not r.correct]
    assert not bad
    assert len(res) == 16


def test_blank_workbook_not_attempted(files):
    practice, _, _ = files
    assert all(r.status == "not_attempted" for r in wbk.grade_workbook(practice))


def test_mistakes_get_specific_feedback(files):
    practice, _, d = files
    wb = openpyxl.load_workbook(practice)
    e, rd = wb["Exercises"], wb["Raw Data"]
    exp = wbk.expected_values()
    e["D5"] = "=SUM('Raw Data'!E2:E31)"              # misses a row
    e["D6"] = float(exp[2])                          # hard-coded number
    e["D7"] = "=SUM('Raw Data'!D2:D32)"
    e["D8"] = "=SUM('Raw Data'!E2:E32)/D7*100"       # percent x100
    e["D9"] = "=D6/SUM('Raw Data'!E2:E32)"                            # right value, missing IFERROR
    e["D10"] = "=D6/0"                               # #DIV/0!
    for row in range(2, 33):                         # fill-down without $
        rd[f"I{row}"] = f"=F{row}/SUM(F{row}:F{row + 30})"
    p = d / "student.xlsx"
    wb.save(p)
    r = {x.task.id: x for x in wbk.grade_workbook(p)}
    assert r[1].status == "incorrect" and "range covers" in r[1].feedback
    assert r[2].status == "hardcoded"
    assert r[3].status == "correct"
    assert r[4].status == "incorrect" and "multiplied by 100" in r[4].feedback
    assert r[5].status == "partial" and "IFERROR" in r[5].feedback
    assert r[6].status == "incorrect" and "#DIV/0!" in r[6].feedback
    assert r[16].status == "incorrect" and "$" in r[16].feedback


def test_lab_evaluation():
    eng = wbk.lab_engine()
    t4 = wbk.TASKS[3]
    assert wbk.evaluate_lab_formula(t4, "D5/D7", eng).correct
    assert not wbk.evaluate_lab_formula(t4, "=D5/D6", eng).correct
    t16 = wbk.TASKS[15]
    assert wbk.evaluate_lab_formula(t16, "=F2/SUM($F$2:$F$32)", eng).correct
    assert not wbk.evaluate_lab_formula(t16, "=F2/SUM(F2:F32)", eng).correct
    t10 = wbk.TASKS[9]
    assert wbk.evaluate_lab_formula(t10, '=INDEX(Budget!C2:C13,MATCH("Jan",Budget!A2:A13,0))', eng).correct


def test_bad_file(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_text("not a workbook")
    with pytest.raises(ValueError):
        wbk.grade_workbook(p)
    wb = openpyxl.Workbook()
    wb.save(tmp_path / "y.xlsx")
    with pytest.raises(ValueError, match="sheet"):
        wbk.grade_workbook(tmp_path / "y.xlsx")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_answer_key_formulas_recalculate_in_libreoffice(files, tmp_path):
    """Independent check: a real spreadsheet program calculates the answer-key
    formulas to the same values as the Python answer key."""
    _, key, _ = files
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir",
                    str(out), str(key)], check=True, capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(out / key.name, data_only=True)
    exp = wbk.expected_values()
    e = wb["Exercises"]
    for t in wbk.TASKS:
        if t.fill_range:
            vals = [wb["Raw Data"][f"I{r}"].value for r in range(2, 33)]
            assert vals == [pytest.approx(x, abs=1e-9) for x in exp[16]]
        elif t.unit == "text":
            assert e[t.cell].value == exp[t.id]
        else:
            assert e[t.cell].value == pytest.approx(float(exp[t.id]), rel=1e-9), t.id
